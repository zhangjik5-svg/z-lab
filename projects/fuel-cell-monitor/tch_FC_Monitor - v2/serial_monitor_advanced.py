# -*- coding: utf-8 -*-
"""tch_FC_Monitor 高级监控版。

本文件在 ``serial_monitor.py`` 的完整保存版基础上扩展，不修改已有两个版本。
新增功能：
1. 参数越限告警与故障码识别；
2. 串口异常后的自动重连；
3. 异常帧、LOG 请求/完整响应和丢弃字节统计；
4. 累计输出电量、输入功率存在时的效率统计；
5. 历史 xlsx 数据加载和回放；
6. 收发/告警/异常日志落盘；
7. 停止采集时自动生成文本分析报告；
8. 按协议进行提示符、帧结构和关键字段完整性校验。

说明：协议给出了 BattSoc，但注明功能未启用时返回 -1；本版只接受 0~100 的设备
上报值，不用电压伪造 SOC。协议没有燃料输入功率，效率仅在未来收到 PIn，或同时
收到 UIn/IIn 时计算。协议 1.0.0 未定义 CRC/校验和，因此不擅自套用校验算法。
"""

import os
import re
import threading
from collections import deque
from datetime import datetime
from queue import Empty

import serial
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from serial_monitor import (
    HAS_OPENPYXL,
    DataRecorder,
    SerialMonitorApp,
    _to_float,
    messagebox,
    tk,
    ttk,
)

try:
    from tkinter import filedialog
except ImportError:  # pragma: no cover - 主程序启动时 tkinter 已检查
    filedialog = None


# 告警阈值是通用初值。投入正式试验前，应按电堆和控制器规格书调整。
# 格式：协议字段 -> (下限, 上限, 单位, 中文名称)
ALARM_LIMITS = {
    "POut": (0.0, 200.0, "W", "输出功率"),
    "UBat": (20.0, 32.0, "V", "电池电压"),
    "UOut": (0.0, 35.0, "V", "输出电压"),
    "IOut": (0.0, 10.0, "A", "输出电流"),
    "TAmb": (-20.0, 50.0, "℃", "环境温度"),
    "TStack": (-20.0, 85.0, "℃", "电堆温度"),
}

FAULT_KEYS = (
    "ActiveErrors",  # LOG 命令中协议规定的当前错误估值码
    "FaultCode",
    "Fault",
    "ErrorCode",
    "AlarmCode",
)
NORMAL_FAULT_VALUES = {"", "0", "0X0", "OK", "NONE", "NORMAL", "NO"}


def _is_zero_code(raw):
    """判断十进制/十六进制故障码是否为零，避免把 0X01 误判为零。"""
    token = str(raw).strip().upper()
    if token == "0X":  # 协议 WARNING 示例用 0X 表示没有错误
        return True
    if re.fullmatch(r"0X[0-9A-F]+", token):
        return int(token[2:], 16) == 0
    value = _to_float(token)
    return value == 0


class ProtocolDataRecorder(DataRecorder):
    """按协议保存 LOG 命令返回的全部字段，而不仅是界面上的曲线字段。"""

    # 第三个元素为 True 时去除 V/A/W/C/%/h 等单位后保存为 Excel 数值。
    # 状态码、原因码和时间保留协议原文，以免数值转换丢失语义。
    FIELD_SPECS = [
        ("输出功率(W)", "POut", True),
        ("电池电压(V)", "UBat", True),
        ("输出电压(V)", "UOut", True),
        ("输出电流(A)", "IOut", True),
        ("环境温度(℃)", "TAmb", True),
        ("电堆温度(℃)", "TStack", True),
        ("热交换器温度(℃)", "THE", True),
        ("甲醇温度(℃)", "TMeOH", True),
        ("空气入口压差", "pDifComp", True),
        ("环境压力(kPa)", "pAmb", True),
        ("相对湿度(%)", "RH", True),
        ("燃料余量(%)", "FL", True),
        ("剩余燃料总量", "MeOHTotal", True),
        ("内部值DSV", "DSV", False),
        ("累计加注时间", "FilltimeTotal", True),
        ("系统上电时间(h)", "SysOpTime", True),
        ("电堆运行时间(h)", "StackOpTime", True),
        ("系统启动次数", "SystemStarts", True),
        ("电堆启动次数", "StackStarts", True),
        ("系统累计功率WOutCum", "WOutCum", True),
        ("系统当前状态", "SystemState", False),
        ("内部状态DmfcState", "DmfcState", False),
        ("燃料电池运行阶段", "DmfcPhase", False),
        ("系统开启原因", "SystemOn", False),
        ("系统关闭原因", "SystemOff", False),
        ("允许手动打开", "SystemOnOK", False),
        ("允许手动关闭", "SystemOffOk", False),
        ("充电模式", "AutoCharge", False),
        ("电池SOC(%)", "BattSoc", True),
        ("电堆最低温度(℃)", "TStackMin", True),
        ("最低温度发生时间", "TStackMinTime", False),
        ("电堆最高温度(℃)", "TStackMax", True),
        ("最高温度发生时间", "TStackMaxTime", False),
        ("当前错误估值码", "ActiveErrors", False),
        ("设备系统时间", "SystemTime", False),
    ]
    COLUMNS = ["序号", "电脑采集时间"] + [item[0] for item in FIELD_SPECS]
    FIELDS = [item[1] for item in FIELD_SPECS]

    def _new_file(self):
        super()._new_file()
        # 基础版只有 12 列；高级版按实际字段数动态设置筛选区域和列宽。
        last_column = get_column_letter(len(self.COLUMNS))
        self.ws.auto_filter.ref = f"A1:{last_column}1"
        for index in range(3, len(self.COLUMNS) + 1):
            self.ws.column_dimensions[get_column_letter(index)].width = 19
        self._save()


class FileLogger:
    """把重要运行信息立即写入日志，程序异常退出时也尽量保留现场。"""

    def __init__(self, log_dir):
        os.makedirs(log_dir, exist_ok=True)
        filename = datetime.now().strftime("monitor_%Y_%m_%d_%H_%M_%S.log")
        self.filepath = os.path.join(log_dir, filename)
        # 行缓冲：每写完一行就交给操作系统，减少意外退出时的日志损失。
        self._file = open(self.filepath, "a", encoding="utf-8", buffering=1)

    def write(self, category, message):
        if self._file is None:
            return
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        clean = str(message).replace("\r", "\\r").replace("\n", "\\n")
        self._file.write(f"{timestamp} [{category}] {clean}\n")

    def close(self):
        if self._file is not None:
            self._file.flush()
            self._file.close()
            self._file = None


class FrameValidator:
    """验证一帧数据是否具备可用于监控和保存的基本结构。"""

    REQUIRED_FIELDS = ("POut", "UOut", "IOut", "TAmb", "TStack")

    # 协议 1.0.0 未说明校验算法，所以默认关闭。仅当后续协议明确采用 XOR8
    # 且覆盖范围一致时，才可改为 "xor8"。
    CHECKSUM_MODE = None
    CHECKSUM_PATTERN = re.compile(
        r"\b(?:Checksum|CRC)\s+(?:0x)?([0-9A-Fa-f]{1,2})\b"
    )

    def validate(self, parsed, payload):
        problems = []
        for field in self.REQUIRED_FIELDS:
            if field not in parsed:
                problems.append(f"缺少字段 {field}")
            elif _to_float(parsed[field]) is None:
                problems.append(f"字段 {field} 不是有效数值")

        checksum_result = self._validate_checksum(payload)
        if checksum_result is False:
            problems.append("XOR8 校验和不匹配")
        return not problems, problems

    def _validate_checksum(self, payload):
        """返回 True/False 表示校验结果，None 表示协议未启用校验。"""
        if self.CHECKSUM_MODE != "xor8":
            return None
        match = self.CHECKSUM_PATTERN.search(payload)
        if match is None:
            return False
        expected = int(match.group(1), 16)
        actual = 0
        # 按校验字段之前的 UTF-8 字节逐字节异或；具体覆盖范围需以协议为准。
        for byte in payload[:match.start()].rstrip().encode("utf-8"):
            actual ^= byte
        return actual == expected


class AlarmManager:
    """检测越限和故障状态；同一故障持续存在时只告警一次。"""

    def __init__(self):
        self.active = {}

    def check(self, parsed):
        new_events = []

        for field, (low, high, unit, label) in ALARM_LIMITS.items():
            value = _to_float(parsed.get(field))
            if value is None:
                continue
            low_key = f"limit:{field}:low"
            high_key = f"limit:{field}:high"
            if value < low:
                message = f"{label}过低：{value:g}{unit} < {low:g}{unit}"
                self._activate(low_key, message, new_events)
                self.active.pop(high_key, None)
            elif value > high:
                message = f"{label}过高：{value:g}{unit} > {high:g}{unit}"
                self._activate(high_key, message, new_events)
                self.active.pop(low_key, None)
            else:
                self.active.pop(low_key, None)
                self.active.pop(high_key, None)

        # 支持几种常见故障字段名；设备实际返回其中任意一个都能识别。
        for key in FAULT_KEYS:
            if key not in parsed:
                continue
            value = str(parsed[key]).strip()
            alarm_key = f"fault:{key}"
            if value.upper() in NORMAL_FAULT_VALUES or _is_zero_code(value):
                self.active.pop(alarm_key, None)
            else:
                message = f"设备故障 {key}={value}"
                self._activate(alarm_key, message, new_events)

        return new_events

    def _activate(self, key, message, events):
        # 同一种异常持续存在时只更新当前值，不重复生成事件；恢复后再次
        # 越限才会重新进入这里并触发新告警。
        if key not in self.active:
            events.append(message)
        self.active[key] = message

    def summary(self):
        return "；".join(self.active.values()) if self.active else "正常"


class SessionStatistics:
    """记录当前采集会话的质量指标和测量值统计。"""

    MEASURE_FIELDS = ("POut", "UBat", "UOut", "IOut", "TAmb", "TStack")
    FIELD_LABELS = {
        "POut": "输出功率(W)",
        "UBat": "电池电压(V)",
        "UOut": "输出电压(V)",
        "IOut": "输出电流(A)",
        "TAmb": "环境温度(℃)",
        "TStack": "电堆温度(℃)",
    }

    def __init__(self):
        self.reset()

    def reset(self):
        self.started_at = datetime.now()
        self.ended_at = None
        self.frames = 0
        self.invalid_frames = 0
        self.lost_frames = 0
        self.sequence_seen = False
        self.log_requests = 0
        self.log_responses = 0
        self.discarded_bytes = 0
        self.reconnects = 0
        self.samples = 0
        self.energy_wh = 0.0
        self.last_sample_time = None
        self.last_power = None
        self.last_sequence = None
        self.field_stats = {}
        self.efficiency_sum = 0.0
        self.efficiency_count = 0
        self.last_batt_soc = None
        self.alarm_events = []
        self.error_history = []

    def observe_frame(self, parsed, valid):
        self.frames += 1
        if not valid:
            self.invalid_frames += 1

        # 若协议带 Seq 或 Sequence，才能可靠推断传输途中丢了多少帧。
        seq_value = parsed.get("Seq", parsed.get("Sequence"))
        seq = _to_float(seq_value)
        if seq is not None:
            self.sequence_seen = True
            seq = int(seq)
            if self.last_sequence is not None and seq > self.last_sequence + 1:
                self.lost_frames += seq - self.last_sequence - 1
            self.last_sequence = seq

    def observe_sample(self, parsed, now=None):
        now = now or datetime.now()
        self.samples += 1
        for field in self.MEASURE_FIELDS:
            value = _to_float(parsed.get(field))
            if value is None:
                continue
            item = self.field_stats.setdefault(
                field, {"count": 0, "sum": 0.0, "min": value, "max": value}
            )
            item["count"] += 1
            item["sum"] += value
            item["min"] = min(item["min"], value)
            item["max"] = max(item["max"], value)

        power = _to_float(parsed.get("POut"))
        if power is not None and self.last_power is not None and self.last_sample_time:
            seconds = (now - self.last_sample_time).total_seconds()
            # 超过一分钟通常意味着串口中断；不把断线时间误算成持续发电。
            if 0 < seconds <= 60:
                self.energy_wh += (self.last_power + power) / 2.0 * seconds / 3600.0
        if power is not None:
            self.last_power = power
            self.last_sample_time = now

        # 只有协议给出输入功率，效率才有物理意义。
        input_power = _to_float(parsed.get("PIn"))
        if input_power is None:
            u_in = _to_float(parsed.get("UIn"))
            i_in = _to_float(parsed.get("IIn"))
            if u_in is not None and i_in is not None:
                input_power = u_in * i_in
        if power is not None and input_power is not None and input_power > 0:
            efficiency = power / input_power * 100.0
            self.efficiency_sum += efficiency
            self.efficiency_count += 1

        batt_soc = _to_float(parsed.get("BattSoc"))
        # 协议说明功能未启用时返回 -1，因此只有 0~100 才作为有效 SOC。
        if batt_soc is not None and 0 <= batt_soc <= 100:
            self.last_batt_soc = batt_soc

    def add_alarm(self, message):
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.alarm_events.append(f"{timestamp} {message}")

    @property
    def average_efficiency(self):
        if not self.efficiency_count:
            return None
        return self.efficiency_sum / self.efficiency_count

    @property
    def unanswered_log_requests(self):
        return max(0, self.log_requests - self.log_responses)


class ReportGenerator:
    """生成便于直接查看和归档的 UTF-8 文本分析报告。"""

    @staticmethod
    def generate(stats, data_dir, data_file=None, checksum_mode=None):
        os.makedirs(data_dir, exist_ok=True)
        base_name = datetime.now().strftime("运行分析报告_%Y_%m_%d_%H_%M_%S")
        path = os.path.join(data_dir, base_name + ".txt")
        suffix = 1
        while os.path.exists(path):
            path = os.path.join(data_dir, f"{base_name}_{suffix:02d}.txt")
            suffix += 1

        stats.ended_at = datetime.now()
        lines = [
            "燃料电池监控运行分析报告",
            "=" * 36,
            f"开始时间：{stats.started_at:%Y-%m-%d %H:%M:%S}",
            f"结束时间：{stats.ended_at:%Y-%m-%d %H:%M:%S}",
            f"数据文件：{data_file or '无'}",
            "",
            "数据质量",
            f"  LOG请求数：{stats.log_requests}",
            f"  LOG完整响应数：{stats.log_responses}",
            f"  未收到对应响应：{stats.unanswered_log_requests}",
            f"  接收完整帧：{stats.frames}",
            f"  有效采样数：{stats.samples}",
            f"  异常帧数：{stats.invalid_frames}",
            (
                f"  序号推断丢帧：{stats.lost_frames}"
                if stats.sequence_seen
                else "  序号推断丢帧：不可用（LOG协议无帧序号）"
            ),
            f"  丢弃无效字节：{stats.discarded_bytes}",
            f"  自动重连次数：{stats.reconnects}",
            f"  校验模式：{checksum_mode or '结构/关键字段校验（校验和未启用）'}",
            "",
            "运行计算",
            f"  累计输出电量：{stats.energy_wh:.4f} Wh",
        ]
        efficiency = stats.average_efficiency
        if efficiency is None:
            lines.append("  平均效率：无法计算（协议未提供 PIn 或 UIn/IIn）")
        else:
            lines.append(f"  平均效率：{efficiency:.2f}%")
        if stats.last_batt_soc is None:
            lines.append("  电池SOC：不可用（协议 BattSoc 返回 -1 或本次未收到有效值）")
        else:
            lines.append(f"  电池SOC：{stats.last_batt_soc:.1f}%（设备 BattSoc）")

        lines.extend(["", "测量统计（最小值 / 平均值 / 最大值）"])
        for field in SessionStatistics.MEASURE_FIELDS:
            item = stats.field_stats.get(field)
            if not item:
                lines.append(f"  {SessionStatistics.FIELD_LABELS[field]}：无数据")
                continue
            average = item["sum"] / item["count"]
            lines.append(
                f"  {SessionStatistics.FIELD_LABELS[field]}："
                f"{item['min']:.4g} / {average:.4g} / {item['max']:.4g}"
            )

        lines.extend(["", f"告警事件：{len(stats.alarm_events)}"])
        if stats.alarm_events:
            lines.extend(f"  {event}" for event in stats.alarm_events)
        else:
            lines.append("  无")

        lines.extend(["", f"设备历史错误码：{len(stats.error_history)}"])
        if stats.error_history:
            lines.extend(f"  {item}" for item in stats.error_history)
        else:
            lines.append("  未查询或设备未返回错误记录")

        with open(path, "w", encoding="utf-8-sig") as report:
            report.write("\n".join(lines) + "\n")
        return path


class AdvancedSerialMonitorApp(SerialMonitorApp):
    """在基础监控界面上增加可靠性、统计、告警和历史回放功能。"""

    DISPLAY_FIELDS = SerialMonitorApp.DISPLAY_FIELDS + [
        ("系统状态", "SystemState"),
        ("运行阶段", "DmfcPhase"),
        ("燃料余量", "FL"),
        ("电池SOC", "BattSoc"),
        ("累计功率", "WOutCum"),
        ("当前错误", "ActiveErrors"),
    ]
    COMMAND_PROMPT = "PFC>"
    MAX_RECV_BUFFER = 1024 * 1024
    RECONNECT_DELAY_MS = 3000

    def __init__(self, root):
        base_dir = os.path.dirname(os.path.abspath(__file__))
        self.file_logger = FileLogger(os.path.join(base_dir, "logs"))
        self.validator = FrameValidator()
        self.alarms = AlarmManager()
        self.stats = SessionStatistics()

        self.connection_requested = False
        self.reconnect_after_id = None
        self.history_rows = []
        self.history_index = 0
        self.history_running = False
        self.history_after_id = None
        self.last_report_path = None
        self.current_data_file = None
        # 每次成功发送的命令依次入队，收到下一个 PFC> 提示符时与响应配对。
        self.pending_commands = deque()
        super().__init__(root)
        self.root.title("tch_FC_Monitor 高级监控版")
        self.root.geometry("1320x900")
        self.file_logger.write("SYSTEM", "高级监控程序启动")

    def _build_ui(self):
        super()._build_ui()

        # 高级功能条独立放在窗口底部，不改变基础版左右布局。
        advanced = ttk.LabelFrame(self.root, text="高级监控")
        advanced.pack(side="bottom", fill="x", padx=8, pady=(2, 8))

        ttk.Button(advanced, text="加载历史", command=self._load_history).pack(
            side="left", padx=5, pady=5
        )
        self.history_btn = ttk.Button(
            advanced, text="播放历史", command=self._toggle_history
        )
        self.history_btn.pack(side="left", padx=5, pady=5)
        ttk.Label(advanced, text="回放速度:").pack(side="left", padx=(8, 2))
        self.history_speed_var = tk.StringVar(value="10")
        ttk.Combobox(
            advanced,
            textvariable=self.history_speed_var,
            values=("1", "5", "10", "20", "50"),
            width=4,
            state="readonly",
        ).pack(side="left", padx=2)
        ttk.Button(advanced, text="生成报告", command=self._manual_report).pack(
            side="left", padx=8, pady=5
        )

        self.stats_var = tk.StringVar(value="帧 0 | 异常 0 | 丢帧 0 | 电量 0.000 Wh")
        ttk.Label(advanced, textvariable=self.stats_var).pack(side="left", padx=12)
        self.alarm_var = tk.StringVar(value="告警：正常")
        ttk.Label(advanced, textvariable=self.alarm_var, foreground="#C00000").pack(
            side="right", padx=10
        )

    # ---------------- 自动重连 ----------------
    def _toggle_port(self):
        # 重连等待状态下再次点击按钮，含义是取消重连并关闭。
        if self.serial_open or self.connection_requested:
            self._close_port()
        else:
            self._open_port()

    def _open_port(self, reconnecting=False):
        port = self.port_var.get().strip()
        if not port:
            if not reconnecting:
                messagebox.showwarning("提示", "请选择或输入串口")
            return
        try:
            baud = int(self.baud_var.get())
        except ValueError:
            if not reconnecting:
                messagebox.showwarning("提示", "波特率无效")
            return

        self.connection_requested = True
        try:
            self.ser = serial.Serial(
                port, baudrate=baud, timeout=0.1, write_timeout=1.0
            )
        except Exception as exc:
            self.ser = None
            self.serial_open = False
            self.open_btn.config(text="取消重连")
            self.status_var.set(f"连接失败，3秒后重试：{exc}")
            self.file_logger.write("SERIAL", f"连接 {port} 失败：{exc}")
            if not reconnecting:
                messagebox.showwarning("连接失败", f"无法打开 {port}，将自动重试：\n{exc}")
            self._schedule_reconnect()
            return

        self.serial_open = True
        self.recv_buffer_str = ""
        self.open_btn.config(text="关闭串口")
        self.status_var.set(f"已打开 {port} @ {baud}")
        self.read_running = True
        self.read_thread = threading.Thread(target=self._read_loop, daemon=True)
        self.read_thread.start()
        if reconnecting:
            self.stats.reconnects += 1
            self.file_logger.write("SERIAL", f"自动重连成功：{port} @ {baud}")
        else:
            self.file_logger.write("SERIAL", f"打开串口：{port} @ {baud}")

    def _schedule_reconnect(self):
        if self.connection_requested and self.reconnect_after_id is None:
            self.reconnect_after_id = self.root.after(
                self.RECONNECT_DELAY_MS, self._attempt_reconnect
            )

    def _attempt_reconnect(self):
        self.reconnect_after_id = None
        if self.connection_requested and not self.serial_open:
            self._open_port(reconnecting=True)

    def _handle_serial_error(self, message):
        if not self.connection_requested:
            return
        self.read_running = False
        if self.ser is not None:
            try:
                self.ser.close()
            except Exception:
                pass
        self.ser = None
        self.serial_open = False
        self.open_btn.config(text="取消重连")
        self.status_var.set("串口中断，正在等待自动重连")
        self.file_logger.write("SERIAL_ERROR", message)
        # 断线前已发出但未完成的命令计入未响应，不再与重连后的响应错配。
        self.pending_commands.clear()
        self._schedule_reconnect()

    def _close_port(self):
        self.connection_requested = False
        if self.reconnect_after_id is not None:
            self.root.after_cancel(self.reconnect_after_id)
            self.reconnect_after_id = None
        super()._close_port()
        self.file_logger.write("SERIAL", "串口关闭/重连取消")

    def _read_loop(self):
        # 使用局部串口对象，防止旧读取线程误读重连后创建的新对象。
        ser_obj = self.ser
        while self.read_running and ser_obj is not None and self.ser is ser_obj:
            try:
                data = ser_obj.read(4096)
            except Exception as exc:
                self.recv_queue.put(("serial_error", str(exc)))
                return
            if data:
                self.recv_queue.put(data)

    def _send(self, msg):
        if not self.serial_open or self.ser is None:
            return False
        try:
            self.ser.write((msg + "\r\n").encode("utf-8"))
        except Exception as exc:
            self._log_send(f"[发送失败: {exc}]")
            self.recv_queue.put(("serial_error", f"发送失败：{exc}"))
            return False
        self._log_send(msg)
        command = msg.strip().split(maxsplit=1)[0].upper() if msg.strip() else ""
        if command:
            self.pending_commands.append(command)
            if command == self.LOOP_CMD:
                self.stats.log_requests += 1
        return True

    def _poll_recv_queue(self):
        try:
            while True:
                item = self.recv_queue.get_nowait()
                if isinstance(item, tuple) and item[0] == "serial_error":
                    self._handle_serial_error(item[1])
                else:
                    self._handle_recv(item)
        except Empty:
            pass
        self.root.after(80, self._poll_recv_queue)

    # ---------------- 帧解析、校验、告警 ----------------
    def _handle_recv(self, data):
        text = data.decode("utf-8", errors="replace")
        self.recv_buffer_str += text

        # 协议以 PFC> 作为每次命令响应结束、下一次命令输入开始的提示符。
        # 缓冲区上限防止设备长期不返回提示符时占满内存。
        if len(self.recv_buffer_str) > self.MAX_RECV_BUFFER:
            discarded = len(self.recv_buffer_str) - self.MAX_RECV_BUFFER
            self.stats.discarded_bytes += discarded
            self.recv_buffer_str = self.recv_buffer_str[-self.MAX_RECV_BUFFER:]

        while self.COMMAND_PROMPT in self.recv_buffer_str:
            end_index = self.recv_buffer_str.find(self.COMMAND_PROMPT)
            response = self.recv_buffer_str[:end_index]
            self.recv_buffer_str = self.recv_buffer_str[
                end_index + len(self.COMMAND_PROMPT):
            ]
            if not response.strip():
                # 打开串口后设备可能先单独发一个 PFC>，它不是命令响应。
                continue

            command = self.pending_commands.popleft() if self.pending_commands else ""
            body = response.strip()
            if not command:
                match = re.match(r"([A-Z]+)\b", body, re.IGNORECASE)
                if match:
                    command = match.group(1).upper()

            # 实机的 LOG 回显可能与首字段直接粘连成 "LOGPOut"，没有
            # 空格或换行。因此已知当前等待 LOG 响应时，直接去掉开头的
            # LOG；其它命令仍要求单词边界，避免把 Information 等正常
            # 响应文本误认为 INFO 命令回显。
            if command == "LOG" and body.upper().startswith("LOG"):
                body = body[len(command):].lstrip()
            elif command and re.match(
                rf"^{re.escape(command)}\b", body, re.IGNORECASE
            ):
                body = body[len(command):].lstrip()

            self._log_recv(response + self.COMMAND_PROMPT)
            self._handle_command_response(command, body)

    def _handle_command_response(self, command, body):
        """根据真实协议分别处理 LOG、WARNING、ERROR 等命令响应。"""
        command = command.upper()
        if command == "LOG":
            self.stats.log_responses += 1
            self._parse_recv_line(body)
        elif command == "WARNING":
            self._parse_warning_response(body)
        elif command in ("ERROR", "ERRORSTAT"):
            self._parse_error_response(command, body)
        else:
            # BAT/INFO/RTC 等命令不参与实时曲线，但完整内容已经显示并落盘。
            self.file_logger.write(f"RESPONSE_{command or 'UNKNOWN'}", body)

    def _parse_warning_response(self, body):
        """解析协议示例中的 ErrorCode1:0X ErrorCode2:0X 格式。"""
        matches = re.findall(r"\b(ErrorCode\d+)\s*:\s*(\S+)", body)
        for key, value in matches:
            alarm_key = f"warning:{key}"
            if _is_zero_code(value) or value.upper() in NORMAL_FAULT_VALUES:
                self.alarms.active.pop(alarm_key, None)
                continue
            message = f"设备警告 {key}={value}"
            if alarm_key not in self.alarms.active:
                self.stats.add_alarm(message)
                self.file_logger.write("ALARM", message)
            self.alarms.active[alarm_key] = message
        self._refresh_stats_ui()

    def _parse_error_response(self, command, body):
        """保存 ERROR/ERRORSTAT 返回的主错误号.次错误号，不误当作当前告警。"""
        codes = re.findall(r"\b\d{3}\.\d{3}\b", body)
        for code in codes:
            entry = f"{command} {code}"
            if entry not in self.stats.error_history:
                self.stats.error_history.append(entry)
        self.file_logger.write(f"RESPONSE_{command}", body)

    def _parse_recv_line(self, line):
        if not line.strip():
            return
        parsed = {
            match.group(1): match.group(2)
            for match in re.finditer(r"(\w+)\s+(\S+)", line)
        }
        # 双重保险：若某种特殊分包仍让 LOG 与 POut 一起进入字段解析，
        # 将实机字段名 LOGPOut 归一化为协议定义的 POut。
        if "POut" not in parsed and "LOGPOut" in parsed:
            parsed["POut"] = parsed.pop("LOGPOut")
        valid, problems = self.validator.validate(parsed, line)
        self.stats.observe_frame(parsed, valid)

        for display_name, proto_key in self.DISPLAY_FIELDS:
            if proto_key in parsed:
                self.display_vars[display_name].set(parsed[proto_key])
        # 若设备使用 Fault/ErrorCode 等名称，也映射到统一故障显示框。
        for fault_key in FAULT_KEYS:
            if fault_key in parsed:
                self.display_vars["当前错误"].set(parsed[fault_key])
                break

        if self.curve1 is not None:
            self.curve1.add_data(parsed)
        if self.curve2 is not None:
            self.curve2.add_data(parsed)
        if self.recorder is not None:
            self.recorder.add_row(parsed)

        if any(field in parsed for field in ProtocolDataRecorder.FIELDS):
            self.stats.observe_sample(parsed)

        if not valid:
            self.file_logger.write("INVALID_FRAME", "；".join(problems))
        new_alarm_events = self.alarms.check(parsed)
        for event in new_alarm_events:
            self.stats.add_alarm(event)
            self.file_logger.write("ALARM", event)
        # LOG 的 ActiveErrors 只给当前错误估值码。首次发现非零值时，
        # 自动执行只读 WARNING 命令获取协议定义的 ErrorCode1/2 详情。
        if any("ActiveErrors" in event for event in new_alarm_events):
            self._send("WARNING")
        self._refresh_stats_ui()

    def _refresh_stats_ui(self):
        if not hasattr(self, "stats_var"):
            return
        self.stats_var.set(
            f"帧 {self.stats.frames} | 异常 {self.stats.invalid_frames} | "
            f"未响应 {self.stats.unanswered_log_requests} | "
            f"电量 {self.stats.energy_wh:.3f} Wh"
        )
        self.alarm_var.set(f"告警：{self.alarms.summary()}")

    # ---------------- 会话记录与报告 ----------------
    def _toggle_loop(self):
        if self.loop_running:
            self._stop_loop()
            return
        if not self.serial_open:
            messagebox.showinfo("提示", "请先打开串口再循环发送")
            return
        if not HAS_OPENPYXL:
            messagebox.showwarning("缺少依赖", "未安装 openpyxl，无法记录数据")

        self.stats.reset()
        self.alarms.active.clear()
        self._refresh_stats_ui()
        self.loop_running = True
        self.loop_btn.config(text="停止")
        if HAS_OPENPYXL:
            # 高级版使用协议完整记录器，保存 LOG 定义的全部 35 个字段。
            self.recorder = ProtocolDataRecorder(self.data_dir)
            self.recorder.start()
            self.current_data_file = self.recorder.filepath
            self.file_logger.write("RECORD", f"开始记录：{self.current_data_file}")
        self._loop_send()

    def _stop_loop(self):
        was_running = self.loop_running
        if self.recorder is not None:
            self.current_data_file = self.recorder.filepath
        super()._stop_loop()
        if was_running:
            self.file_logger.write("RECORD", "停止记录")
            self._generate_report(show_message=False)

    def _generate_report(self, show_message=False):
        try:
            self.last_report_path = ReportGenerator.generate(
                self.stats,
                self.data_dir,
                self.current_data_file,
                self.validator.CHECKSUM_MODE,
            )
            self.status_var.set(f"报告已生成：{os.path.basename(self.last_report_path)}")
            self.file_logger.write("REPORT", self.last_report_path)
            if show_message:
                messagebox.showinfo("报告完成", self.last_report_path)
        except Exception as exc:
            self.file_logger.write("REPORT_ERROR", str(exc))
            if show_message:
                messagebox.showerror("报告失败", str(exc))

    def _manual_report(self):
        self._generate_report(show_message=True)

    # ---------------- 历史数据加载和回放 ----------------
    def _load_history(self):
        if self.loop_running:
            messagebox.showinfo("提示", "请先停止实时循环采集，再加载历史数据")
            return
        if filedialog is None or not HAS_OPENPYXL:
            messagebox.showerror("不可用", "当前环境无法读取 xlsx 文件")
            return
        path = filedialog.askopenfilename(
            title="选择燃料电池历史数据",
            initialdir=self.data_dir,
            filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        if not path:
            return
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook["FC_data"] if "FC_data" in workbook.sheetnames else workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows)
            header_map = self._history_header_map(headers)
            history = []
            for values in rows:
                parsed = {}
                for index, field in header_map.items():
                    if index < len(values) and values[index] not in (None, ""):
                        parsed[field] = values[index]
                if parsed:
                    timestamp = values[1] if len(values) > 1 else None
                    history.append((parsed, timestamp))
            workbook.close()
        except Exception as exc:
            messagebox.showerror("读取失败", str(exc))
            return

        self.history_rows = history
        self.history_index = 0
        self._preview_history()
        self.status_var.set(
            f"已加载 {os.path.basename(path)}，共 {len(history)} 条"
        )
        self.file_logger.write("HISTORY", f"加载 {path}，{len(history)} 条")

    @staticmethod
    def _history_header_map(headers):
        # 同时兼容原版 7 列、完整保存版 12 列和协议高级版 37 列文件。
        aliases = {
            column_name: field
            for column_name, field, _ in ProtocolDataRecorder.FIELD_SPECS
        }
        # 完整保存版使用了不带 (h) 的旧列名。
        aliases["电堆运行时间"] = "StackOpTime"
        return {index: aliases[name] for index, name in enumerate(headers) if name in aliases}

    def _clear_curves(self):
        for curve in (self.curve1, self.curve2):
            if curve is None:
                continue
            for buffer in curve.buffers:
                buffer.clear()
            curve.last_time = None
            curve.redraw()

    def _append_history_point(self, parsed, timestamp):
        # 批量回放时直接追加缓冲，避免每个点都重绘造成界面卡顿。
        for curve in (self.curve1, self.curve2):
            if curve is None:
                continue
            for index, series in enumerate(curve.series):
                curve.buffers[index].append(_to_float(parsed.get(series["field"])))
            if isinstance(timestamp, datetime):
                curve.last_time = timestamp
            else:
                curve.last_time = datetime.now()
        for display_name, proto_key in self.DISPLAY_FIELDS:
            if proto_key in parsed:
                self.display_vars[display_name].set(str(parsed[proto_key]))

    def _preview_history(self):
        self._clear_curves()
        if not self.history_rows:
            return
        max_points = max(self.curve1.max_points, self.curve2.max_points)
        for parsed, timestamp in self.history_rows[-max_points:]:
            self._append_history_point(parsed, timestamp)
        self.curve1.redraw()
        self.curve2.redraw()

    def _toggle_history(self):
        if self.history_running:
            self._stop_history()
            return
        if not self.history_rows:
            messagebox.showinfo("提示", "请先加载历史 Excel 文件")
            return
        if self.loop_running:
            messagebox.showinfo("提示", "请先停止实时循环采集")
            return
        self._clear_curves()
        self.history_index = 0
        self.history_running = True
        self.history_btn.config(text="停止回放")
        self._history_step()

    def _history_step(self):
        if not self.history_running:
            return
        try:
            batch = max(1, int(self.history_speed_var.get()))
        except ValueError:
            batch = 10
        stop = min(len(self.history_rows), self.history_index + batch)
        for index in range(self.history_index, stop):
            self._append_history_point(*self.history_rows[index])
        self.history_index = stop
        self.curve1.redraw()
        self.curve2.redraw()
        self.status_var.set(f"历史回放：{stop}/{len(self.history_rows)}")
        if stop >= len(self.history_rows):
            self._stop_history()
            self.status_var.set("历史回放完成")
        else:
            self.history_after_id = self.root.after(100, self._history_step)

    def _stop_history(self):
        self.history_running = False
        if self.history_after_id is not None:
            self.root.after_cancel(self.history_after_id)
            self.history_after_id = None
        if hasattr(self, "history_btn"):
            self.history_btn.config(text="播放历史")

    # ---------------- 日志与退出 ----------------
    def _log_send(self, msg):
        super()._log_send(msg)
        self.file_logger.write("SEND", msg)

    def _log_recv(self, msg):
        super()._log_recv(msg)
        self.file_logger.write("RECV", msg)

    def _on_close(self):
        self.connection_requested = False
        if self.reconnect_after_id is not None:
            self.root.after_cancel(self.reconnect_after_id)
            self.reconnect_after_id = None
        self._stop_history()
        self.file_logger.write("SYSTEM", "程序退出")
        try:
            super()._on_close()
        finally:
            self.file_logger.close()


def main():
    root = tk.Tk()
    AdvancedSerialMonitorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
